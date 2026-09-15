import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # 16:9 widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6] # blank layout

    # Colors
    BG_DARK = RGBColor(10, 15, 26)        # #0a0f1a
    CARD_BG = RGBColor(17, 24, 39)        # #111827
    CARD_BORDER = RGBColor(30, 41, 59)    # #1e293b
    CYAN_ACCENT = RGBColor(6, 182, 212)   # #06b6d4
    WHITE = RGBColor(241, 245, 249)       # #f1f5f9
    GRAY_TEXT = RGBColor(148, 163, 184)   # #94a3b8
    HIGH_RED = RGBColor(255, 77, 77)      # #ff4d4d
    SAFE_GREEN = RGBColor(40, 167, 69)    # #28a745
    AMBER = RGBColor(245, 158, 11)        # #f59e0b

    def add_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category="RETAIL PULSE PREDICTION  |  AI SALES FORECASTING & INVENTORY INTELLIGENCE"):
        # Category label
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(11.5), Inches(0.4))
        tf_c = cat_box.text_frame
        tf_c.word_wrap = True
        p_c = tf_c.paragraphs[0]
        p_c.text = category.upper()
        p_c.font.size = Pt(11)
        p_c.font.bold = True
        p_c.font.color.rgb = CYAN_ACCENT

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.5), Inches(0.8))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = Pt(26)
        p_t.font.bold = True
        p_t.font.color.rgb = WHITE

    def create_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
        return card

    # ==========================================
    # SLIDE 1: Title Slide
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    add_background(s1)

    # Title & Subtitle box (Centered, prominent)
    tbox = s1.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.733), Inches(2.3))
    tf = tbox.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "RETAIL PULSE PREDICTION"
    p0.font.size = Pt(44)
    p0.font.bold = True
    p0.font.color.rgb = WHITE
    p0.alignment = PP_ALIGN.CENTER

    p1 = tf.add_paragraph()
    p1.text = "AI-Powered Retail Sales Forecasting & Insights"
    p1.font.size = Pt(22)
    p1.font.bold = True
    p1.font.color.rgb = CYAN_ACCENT
    p1.alignment = PP_ALIGN.CENTER
    p1.space_before = Pt(10)

    # Center card for "Presented By" and Academic details
    card_w = Inches(10.5)
    card_left = (Inches(13.333) - card_w) / 2
    c1 = create_card(s1, card_left, Inches(3.3), card_w, Inches(3.6))

    # Inner content box
    tbox_info = s1.shapes.add_textbox(card_left + Inches(0.5), Inches(3.5), card_w - Inches(1.0), Inches(3.2))
    tfi = tbox_info.text_frame
    tfi.word_wrap = True

    # Presented by header
    p_by = tfi.paragraphs[0]
    p_by.text = "── PRESENTED BY ──"
    p_by.font.size = Pt(13)
    p_by.font.bold = True
    p_by.font.color.rgb = AMBER
    p_by.alignment = PP_ALIGN.CENTER

    # 4 Team members
    members = [
        "Prince Umrao",
        "Raunak Kesharwani",
        "Vibhu Kumar",
        "Shubham Thakur"
    ]
    
    # 4 Team members vertically stacked as requested
    members = [
        "Prince Umrao",
        "Raunak Kesharwani",
        "Vibhu Kumar",
        "Shubham Thakur"
    ]
    for idx, m in enumerate(members):
        p_m = tfi.add_paragraph()
        p_m.text = m
        p_m.font.size = Pt(15)
        p_m.font.bold = True
        p_m.font.color.rgb = WHITE
        p_m.alignment = PP_ALIGN.CENTER
        p_m.space_before = Pt(8 if idx == 0 else 4)

    # Decorative separator
    p_sep = tfi.add_paragraph()
    p_sep.text = "✦   ✦   ✦"
    p_sep.font.size = Pt(9)
    p_sep.font.color.rgb = CYAN_ACCENT
    p_sep.alignment = PP_ALIGN.CENTER
    p_sep.space_before = Pt(10)

    # Department
    p_dept = tfi.add_paragraph()
    p_dept.text = "B.Tech – Computer Science Engineering"
    p_dept.font.size = Pt(14)
    p_dept.font.bold = True
    p_dept.font.color.rgb = CYAN_ACCENT
    p_dept.alignment = PP_ALIGN.CENTER
    p_dept.space_before = Pt(8)

    # University
    p_univ = tfi.add_paragraph()
    p_univ.text = "Lamrin Tech Skills University"
    p_univ.font.size = Pt(16)
    p_univ.font.bold = True
    p_univ.font.color.rgb = WHITE
    p_univ.alignment = PP_ALIGN.CENTER
    p_univ.space_before = Pt(4)

    # ==========================================
    # SLIDE 2: Problem Statement & Industry Dilemma
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    add_background(s2)
    add_header(s2, "The Retail Dilemma: The Hidden Cost of Stock-Outs")

    col_w = Inches(3.64)
    gap = Inches(0.4)

    # Card 1: The Problem
    create_card(s2, Inches(0.8), Inches(1.8), col_w, Inches(4.9))
    tb = s2.shapes.add_textbox(Inches(1.0), Inches(2.0), col_w - Inches(0.4), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "⚠️ Empty Shelves Crisis"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = HIGH_RED

    points1 = [
        "Lost Direct Sales: When an item is out of stock, 31% of shoppers switch stores immediately.",
        "Brand Churn: Customer satisfaction drops drastically; repeated stockouts permanently drive buyers to competitors.",
        "Emergency Logistics: Rush ordering and air freight to restock depleted items skyrocket operational expenditure.",
        "Distorted Demand Signals: When products aren't available, true customer demand data is masked and lost."
    ]
    for pt in points1:
        p = tf.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(12)
        p.font.color.rgb = GRAY_TEXT
        p.space_before = Pt(12)

    # Card 2: Traditional Flaw
    create_card(s2, Inches(0.8) + col_w + gap, Inches(1.8), col_w, Inches(4.9))
    tb = s2.shapes.add_textbox(Inches(1.0) + col_w + gap, Inches(2.0), col_w - Inches(0.4), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "⏱️ The Reactive Trap"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = AMBER

    points2 = [
        "Lagging Indicators: Most ERPs and inventory spreadsheets only alert managers AFTER stock hits zero.",
        "Static Thresholds: Fixed reorder points fail to adjust for velocity spikes, weekend sales surges, or seasonality.",
        "Manual Oversight: Managing 50,000+ SKUs manually in spreadsheets leads to human error and overlooked risks.",
        "Overstocking Counter-Effect: To compensate, managers overbuy slow-moving stock, freezing working capital."
    ]
    for pt in points2:
        p = tf.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(12)
        p.font.color.rgb = GRAY_TEXT
        p.space_before = Pt(12)

    # Card 3: The Retail Pulse Solution
    create_card(s2, Inches(0.8) + (col_w + gap)*2, Inches(1.8), col_w, Inches(4.9))
    tb = s2.shapes.add_textbox(Inches(1.0) + (col_w + gap)*2, Inches(2.0), col_w - Inches(0.4), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🚀 Proactive Intelligence"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT

    points3 = [
        "Predictive Modeling: Employs Scikit-Learn ML models trained on sales patterns to anticipate stock depletion days ahead.",
        "Dynamic Runway Analysis: Computes real-time 'Days Left' based on rolling 7-day sales velocities.",
        "Automated Triage: Segregates safe inventory from critical stockout emergencies instantly in bulk.",
        "Actionable Timeliness: Gives procurement teams the exact window needed to restock before shelves empty."
    ]
    for pt in points3:
        p = tf.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(12)
        p.font.color.rgb = GRAY_TEXT
        p.space_before = Pt(12)

    # ==========================================
    # SLIDE 3: System Overview & Core Value Proposition
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    add_background(s3)
    add_header(s3, "Retail Pulse at a Glance: The Complete Value Proposition")

    # 4 KPI / Value blocks across top
    kpi_w = Inches(2.7)
    kpi_gap = Inches(0.31)
    kpis = [
        ("BULK CSV INGESTION", "50,000+ SKUs", "Instantly parses entire catalog rows with auto-detected columns", CYAN_ACCENT),
        ("ML RISK PREDICTION", "< 100ms / Item", "Random Forest / Classifier inference with risk probability scoring", AMBER),
        ("STOCKOUT RUNWAY", "Days-Left Metric", "Dynamic depletion estimation using 7-day rolling sales trends", SAFE_GREEN),
        ("DECISION SPEED", "10x Faster", "Replaces manual auditing with color-coded triage and charts", HIGH_RED)
    ]

    for i, (title, stat, desc, col) in enumerate(kpis):
        left_pos = Inches(0.8) + i * (kpi_w + kpi_gap)
        create_card(s3, left_pos, Inches(1.8), kpi_w, Inches(2.0))
        tb = s3.shapes.add_textbox(left_pos + Inches(0.15), Inches(1.9), kpi_w - Inches(0.3), Inches(1.8))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.text = stat
        p2.font.size = Pt(22)
        p2.font.bold = True
        p2.font.color.rgb = WHITE
        p2.space_before = Pt(4)

        p3 = tf.add_paragraph()
        p3.text = desc
        p3.font.size = Pt(11)
        p3.font.color.rgb = GRAY_TEXT
        p3.space_before = Pt(4)

    # Bottom Two Deep-Dive Cards
    create_card(s3, Inches(0.8), Inches(4.1), Inches(5.7), Inches(2.7))
    tb = s3.shapes.add_textbox(Inches(1.0), Inches(4.25), Inches(5.3), Inches(2.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🎯 What Retail Pulse Solves"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT

    pts_l = [
        "Eliminates Blind Spots: Supermarkets and retail chains sell thousands of SKUs; high-velocity items can vanish in hours.",
        "Prevents Over-Stocking & Under-Stocking: Balances optimal holding costs with zero stock-out downtime.",
        "Empowers Store Operators: No data science expertise needed; intuitive UI shows red/green statuses and top 5 urgent items."
    ]
    for pt in pts_l:
        p = tf.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(12)
        p.font.color.rgb = GRAY_TEXT
        p.space_before = Pt(6)

    create_card(s3, Inches(6.833), Inches(4.1), Inches(5.7), Inches(2.7))
    tb = s3.shapes.add_textbox(Inches(7.033), Inches(4.25), Inches(5.3), Inches(2.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "💡 Operational Paradigm Shift"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = SAFE_GREEN

    pts_r = [
        "From Retrospective to Predictive: Shifts management focus from 'What ran out yesterday?' to 'What will run out on Friday?'.",
        "Unified Workflow: Drag-and-drop .pkl ML model + transaction CSV -> Instant prioritized restock action list.",
        "Lightweight & Zero-Setup: Runs in any modern browser without heavy server infrastructure, ready for edge retail terminals."
    ]
    for pt in pts_r:
        p = tf.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(12)
        p.font.color.rgb = GRAY_TEXT
        p.space_before = Pt(6)

    # ==========================================
    # SLIDE 4: Architecture & End-to-End Pipeline
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    add_background(s4)
    add_header(s4, "System Architecture: End-to-End Data & ML Pipeline")

    steps = [
        ("1. Data Ingestion", "CSV Upload Engine", [
            "Accepts raw transaction logs & inventory datasets",
            "Auto-detects columns (Product, Sales, Stock, 7d Avg)",
            "Supports 50k+ row datasets (e.g. retail_data_cleaned_final.csv)"
        ], CYAN_ACCENT),
        ("2. Feature Processing", "Velocity & Runway Math", [
            "Computes Days Left = Stock / Avg Sales (7d)",
            "Calculates Sales Velocity = Sales / Avg Sales",
            "Evaluates Stock-to-Sales Cover Ratio"
        ], AMBER),
        ("3. Model Inference", "Scikit-Learn ML Core", [
            "Executes Random Forest / Decision Tree classification",
            "Generates Stock-Out Probability (0.0 to 1.0)",
            "Classifies binary risk: Safe (0) vs High Risk (1)"
        ], SAFE_GREEN),
        ("4. Executive Dashboard", "Visual Triage UI", [
            "Displays High-Risk alert counters & risk percentage",
            "Renders Chart.js risk breakdown & days-left charts",
            "Highlights Top 5 critical items needing immediate restock"
        ], HIGH_RED)
    ]

    pipe_w = Inches(2.7)
    pipe_gap = Inches(0.31)
    for i, (title, subtitle, bullets, col) in enumerate(steps):
        left_pos = Inches(0.8) + i * (pipe_w + pipe_gap)
        create_card(s4, left_pos, Inches(1.8), pipe_w, Inches(5.0))

        # Color bar indicator on top of card
        cbar = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, left_pos, Inches(1.8), pipe_w, Inches(0.08))
        cbar.fill.solid()
        cbar.fill.fore_color.rgb = col
        cbar.line.fill.background()

        tb = s4.shapes.add_textbox(left_pos + Inches(0.15), Inches(2.0), pipe_w - Inches(0.3), Inches(4.6))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.size = Pt(15)
        p2.font.bold = True
        p2.font.color.rgb = WHITE
        p2.space_before = Pt(4)

        for b in bullets:
            p_b = tf.add_paragraph()
            p_b.text = "• " + b
            p_b.font.size = Pt(11)
            p_b.font.color.rgb = GRAY_TEXT
            p_b.space_before = Pt(12)

    # ==========================================
    # SLIDE 5: Machine Learning & Mathematical Logic
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    add_background(s5)
    add_header(s5, "The Intelligence Engine: Feature Engineering & ML Formulation")

    # Left Card: Formula & Features
    create_card(s5, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.0))
    tb_l = s5.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.3), Inches(4.6))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True

    p = tf_l.paragraphs[0]
    p.text = "📐 Mathematical Formulation"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT

    p_f1 = tf_l.add_paragraph()
    p_f1.text = "1. Days of Stock Remaining (Runway):"
    p_f1.font.size = Pt(13)
    p_f1.font.bold = True
    p_f1.font.color.rgb = WHITE
    p_f1.space_before = Pt(10)

    p_eq1 = tf_l.add_paragraph()
    p_eq1.text = "   Days Left = Current Stock / max(Avg Sales [7-day], 0.01)"
    p_eq1.font.size = Pt(12)
    p_eq1.font.color.rgb = AMBER
    p_eq1.space_before = Pt(2)

    p_f2 = tf_l.add_paragraph()
    p_f2.text = "2. Sales Velocity Factor:"
    p_f2.font.size = Pt(13)
    p_f2.font.bold = True
    p_f2.font.color.rgb = WHITE
    p_f2.space_before = Pt(10)

    p_eq2 = tf_l.add_paragraph()
    p_eq2.text = "   Velocity = Current Sales / max(Avg Sales [7-day], 0.01)"
    p_eq2.font.size = Pt(12)
    p_eq2.font.color.rgb = AMBER
    p_eq2.space_before = Pt(2)

    p_f3 = tf_l.add_paragraph()
    p_f3.text = "3. Stock-to-Sales Cover Ratio:"
    p_f3.font.size = Pt(13)
    p_f3.font.bold = True
    p_f3.font.color.rgb = WHITE
    p_f3.space_before = Pt(10)

    p_eq3 = tf_l.add_paragraph()
    p_eq3.text = "   Cover Ratio = Current Stock / max(Current Sales, 0.01)"
    p_eq3.font.size = Pt(12)
    p_eq3.font.color.rgb = AMBER
    p_eq3.space_before = Pt(2)

    p_f4 = tf_l.add_paragraph()
    p_f4.text = "Threshold Rules: Days Left <= 3d triggers immediate critical priority; velocity > 1.4x signals rapid demand acceleration."
    p_f4.font.size = Pt(11)
    p_f4.font.color.rgb = GRAY_TEXT
    p_f4.space_before = Pt(12)

    # Right Card: Model Architecture
    create_card(s5, Inches(6.833), Inches(1.8), Inches(5.7), Inches(5.0))
    tb_r = s5.shapes.add_textbox(Inches(7.033), Inches(2.0), Inches(5.3), Inches(4.6))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True

    p = tf_r.paragraphs[0]
    p.text = "🤖 Model Architecture & Classification"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = SAFE_GREEN

    pts_model = [
        "Algorithm: Random Forest & Decision Tree Classifiers trained on historical retail transactions and stockouts.",
        "Target Output (Binary): \n  • Class 0 (Safe): Inventory runway >= 7 days with steady demand.\n  • Class 1 (High Risk): Imminent stock exhaustion within replenishment lead time.",
        "Probability Scoring: Calculates risk confidence score (0% to 100%) by evaluating runway, velocity, and stock coverage.",
        "Model Portability: Exported as standardized .pkl / .joblib artifacts for instant drag-and-drop loading in client or backend API.",
        "Robust Fallbacks: Handles zero-division safeguards and missing headers gracefully with dynamic pattern matching."
    ]
    for pt in pts_model:
        p = tf_r.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(12)
        p.font.color.rgb = GRAY_TEXT
        p.space_before = Pt(8)

    # ==========================================
    # SLIDE 6: Interactive Dashboard & UI Highlights
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    add_background(s6)
    add_header(s6, "Interactive Dashboard: High-Impact Decision Support UI")

    # 3 Column Cards
    col3_w = Inches(3.64)
    c3_gap = Inches(0.4)

    # UI Feature 1: Modern Dark Glassmorphism
    create_card(s6, Inches(0.8), Inches(1.8), col3_w, Inches(5.0))
    tb = s6.shapes.add_textbox(Inches(1.0), Inches(2.0), col3_w - Inches(0.4), Inches(4.6))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🎨 Glassmorphic Dark UI"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT

    pts_ui1 = [
        "Space Grotesk & DM Sans typography with high contrast visibility for warehouse & ops managers.",
        "Collapsible sidebar controls for drag-and-drop model (.pkl) & bulk data (.csv) uploads.",
        "Live status indicators (Connected/Disconnected) and animated atmospheric ambient background.",
        "Responsive layout adapted for desktop monitors and tablet/mobile operational floor inspection."
    ]
    for pt in pts_ui1:
        p = tf.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(12)
        p.font.color.rgb = GRAY_TEXT
        p.space_before = Pt(10)

    # UI Feature 2: Real-time Analytics & Charts
    create_card(s6, Inches(0.8) + col3_w + c3_gap, Inches(1.8), col3_w, Inches(5.0))
    tb = s6.shapes.add_textbox(Inches(1.0) + col3_w + c3_gap, Inches(2.0), col3_w - Inches(0.4), Inches(4.6))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "📊 Dynamic Visualizations"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = AMBER

    pts_ui2 = [
        "Risk Distribution Chart: Instant visual split between safe SKUs and stock-out hazards.",
        "Days-Left Trend Chart: Ranks top 30 critical products by days of remaining inventory.",
        "KPI Metric Counter: Live count of Total Products Analyzed, High Risk Alerts, and Safe Stock.",
        "Overall Risk Rate Meter: Color-graduated percentage gauge showing catalog health."
    ]
    for pt in pts_ui2:
        p = tf.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(12)
        p.font.color.rgb = GRAY_TEXT
        p.space_before = Pt(10)

    # UI Feature 3: Actionable Table & Top 5 Alerts
    create_card(s6, Inches(0.8) + (col3_w + c3_gap)*2, Inches(1.8), col3_w, Inches(5.0))
    tb = s6.shapes.add_textbox(Inches(1.0) + (col3_w + c3_gap)*2, Inches(2.0), col3_w - Inches(0.4), Inches(4.6))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "⚡ Prioritized Risk Table"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = HIGH_RED

    pts_ui3 = [
        "Automated Urgency Sorting: Automatically elevates high-risk items with shortest runways to the top.",
        "TOP 5 Risky Badge: Distinct visual highlighting for the top 5 most vulnerable products.",
        "Detailed Columns: Product Name, Sales, Available Stock, 7-Day Average, and Exact Days Left.",
        "Color-coded Risk Dots: Instant red/green visual triage reducing cognitive fatigue during replenishment planning."
    ]
    for pt in pts_ui3:
        p = tf.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(12)
        p.font.color.rgb = GRAY_TEXT
        p.space_before = Pt(10)

    # ==========================================
    # SLIDE 7: Business Value, ROI & Competitive Edge
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    add_background(s7)
    add_header(s7, "Quantifiable Business Value & ROI Impact")

    # 4 Quadrants
    quad_w = Inches(5.7)
    quad_h = Inches(2.35)
    gap_x = Inches(0.333)
    gap_y = Inches(0.3)

    quads = [
        ("💰 Revenue Preservation", CYAN_ACCENT, [
            "Reduces stock-out lost sales by up to 35% through timely advance purchase orders.",
            "Captures peak seasonal sales that traditional static threshold reorder points routinely miss."
        ]),
        ("📈 Customer Lifetime Value (LTV)", SAFE_GREEN, [
            "Prevents customer abandonment; 70%+ of retail shoppers who encounter repeated stockouts switch retailers.",
            "Protects omnichannel trust and brand loyalty across online storefronts and physical branches."
        ]),
        ("📦 Working Capital Optimization", AMBER, [
            "Eliminates indiscriminate over-buffering; stops stores from locking cash into stagnant products.",
            "Allows precision allocation of procurement budget to high-velocity, high-margin SKUs."
        ]),
        ("⏱️ Operational Labor Efficiency", HIGH_RED, [
            "Saves hours of manual Excel collation and tedious stock auditing every week.",
            "Generates instant restock priorities with one click, streamlining purchasing team workflows."
        ])
    ]

    for i, (qtitle, qcol, qpoints) in enumerate(quads):
        row = i // 2
        col = i % 2
        x = Inches(0.8) + col * (quad_w + gap_x)
        y = Inches(1.8) + row * (quad_h + gap_y)

        create_card(s7, x, y, quad_w, quad_h)
        tb = s7.shapes.add_textbox(x + Inches(0.2), y + Inches(0.15), quad_w - Inches(0.4), quad_h - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = qtitle
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = qcol

        for qp in qpoints:
            p2 = tf.add_paragraph()
            p2.text = "• " + qp
            p2.font.size = Pt(12)
            p2.font.color.rgb = GRAY_TEXT
            p2.space_before = Pt(6)

    # ==========================================
    # SLIDE 8: Technology Stack & Implementation
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    add_background(s8)
    add_header(s8, "Technology Stack & Engineering Blueprint")

    tech_categories = [
        ("AI & Machine Learning", CYAN_ACCENT, [
            "Scikit-Learn (Random Forest & Decision Trees)",
            "Python 3.10 / Pandas (Data Cleaning & Preprocessing)",
            "Pickle / Joblib (.pkl Model Serialization)",
            "Feature Engineering: 7-day Moving Averages, Velocity Ratios"
        ]),
        ("Interactive Frontend", AMBER, [
            "Modern HTML5 & Vanilla JavaScript (Zero bloated framework runtime)",
            "Tailwind CSS (Utility-first dark theme & micro-interactions)",
            "Chart.js (Interactive Bar & Trend Visualizations)",
            "Google Fonts (Space Grotesk & DM Sans typography)"
        ]),
        ("Data Handling & Storage", SAFE_GREEN, [
            "Client-Side CSV Streaming Parser (Fast batch ingestion)",
            "Retail Cleaned Dataset (58,600+ real transactions)",
            "Smart Header Matcher (Auto-resolves sku, sales, stock columns)",
            "Lightweight in-memory state engine"
        ]),
        ("Deployment & Scalability", WHITE, [
            "Serverless GitHub Pages Ready (Immediate web access)",
            "REST API Extensible (/api/predict endpoint architecture)",
            "Docker / Cloud Containerization capability",
            "Cross-platform compatibility (Desktop, Tablet, POS edge)"
        ])
    ]

    for i, (title, col, items) in enumerate(tech_categories):
        left_pos = Inches(0.8) + i * (pipe_w + pipe_gap)
        create_card(s8, left_pos, Inches(1.8), pipe_w, Inches(5.0))

        cbar = s8.shapes.add_shape(MSO_SHAPE.RECTANGLE, left_pos, Inches(1.8), pipe_w, Inches(0.08))
        cbar.fill.solid()
        cbar.fill.fore_color.rgb = col
        cbar.line.fill.background()

        tb = s8.shapes.add_textbox(left_pos + Inches(0.15), Inches(2.0), pipe_w - Inches(0.3), Inches(4.6))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = col

        for item in items:
            p_i = tf.add_paragraph()
            p_i.text = "• " + item
            p_i.font.size = Pt(11)
            p_i.font.color.rgb = GRAY_TEXT
            p_i.space_before = Pt(12)

    # ==========================================
    # SLIDE 9: Future Roadmap & Growth Vision
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    add_background(s9)
    add_header(s9, "Future Roadmap: Scaling to Enterprise Supply Chains")

    roadmaps = [
        ("PHASE 1: CURRENT", "Prototype & Core Engine", [
            "Bulk CSV prediction engine",
            ".pkl model drag-and-drop loading",
            "Interactive dashboard with charts",
            "Top 5 urgency triage table"
        ], CYAN_ACCENT),
        ("PHASE 2: NEAR TERM", "Automated Alerting", [
            "WhatsApp & Slack alert webhooks",
            "Automated Email notifications to buyers",
            "FastAPI backend with database storage",
            "Multi-store / multi-warehouse aggregation"
        ], AMBER),
        ("PHASE 3: MID TERM", "Deep Supply Chain Sync", [
            "Direct ERP integration (SAP, Oracle, Shopify)",
            "Auto-generate Restock Purchase Orders (PO)",
            "Supplier Lead-Time tracking integration",
            "Dynamic safety stock buffer calculator"
        ], SAFE_GREEN),
        ("PHASE 4: LONG TERM", "Advanced Deep Learning", [
            "Prophet / LSTM time-series demand forecasting",
            "External factors (weather, festivals, promotions)",
            "Cannibalization & cross-elasticity modeling",
            "Autonomous self-healing inventory network"
        ], HIGH_RED)
    ]

    for i, (phase, title, bullets, col) in enumerate(roadmaps):
        left_pos = Inches(0.8) + i * (pipe_w + pipe_gap)
        create_card(s9, left_pos, Inches(1.8), pipe_w, Inches(5.0))

        cbar = s9.shapes.add_shape(MSO_SHAPE.RECTANGLE, left_pos, Inches(1.8), pipe_w, Inches(0.08))
        cbar.fill.solid()
        cbar.fill.fore_color.rgb = col
        cbar.line.fill.background()

        tb = s9.shapes.add_textbox(left_pos + Inches(0.15), Inches(2.0), pipe_w - Inches(0.3), Inches(4.6))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = phase
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.text = title
        p2.font.size = Pt(14)
        p2.font.bold = True
        p2.font.color.rgb = WHITE
        p2.space_before = Pt(4)

        for b in bullets:
            p_b = tf.add_paragraph()
            p_b.text = "• " + b
            p_b.font.size = Pt(11)
            p_b.font.color.rgb = GRAY_TEXT
            p_b.space_before = Pt(10)

    # ==========================================
    # SLIDE 10: Conclusion & Summary
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    add_background(s10)

    # Large centerpiece card
    create_card(s10, Inches(1.2), Inches(1.0), Inches(10.933), Inches(5.5))

    tb_c = s10.shapes.add_textbox(Inches(1.6), Inches(1.3), Inches(10.133), Inches(4.9))
    tf_c = tb_c.text_frame
    tf_c.word_wrap = True

    p = tf_c.paragraphs[0]
    p.text = "RETAIL PULSE"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT

    p_quote = tf_c.add_paragraph()
    p_quote.text = '"We are not just predicting stock-outs; we are enabling proactive inventory intelligence."'
    p_quote.font.size = Pt(24)
    p_quote.font.bold = True
    p_quote.font.color.rgb = WHITE
    p_quote.space_before = Pt(12)

    p_div = tf_c.add_paragraph()
    p_div.text = "Key Takeaways:"
    p_div.font.size = Pt(16)
    p_div.font.bold = True
    p_div.font.color.rgb = CYAN_ACCENT
    p_div.space_before = Pt(20)

    takeaways = [
        "Actionable Machine Learning: Transforms overwhelming transaction data into precise, prioritized replenishment decisions.",
        "Zero Friction Adoption: Drag-and-drop model + CSV file processing; instant analysis without complex deployment hurdles.",
        "Clear Business Impact: Eliminates lost sales, prevents stockouts, protects customer loyalty, and frees up tied working capital.",
        "Enterprise-Ready Architecture: Built on modular Scikit-Learn logic, ready for API integration with enterprise ERP systems."
    ]
    for ta in takeaways:
        p_ta = tf_c.add_paragraph()
        p_ta.text = "✔  " + ta
        p_ta.font.size = Pt(13)
        p_ta.font.color.rgb = GRAY_TEXT
        p_ta.space_before = Pt(8)

    p_auth = tf_c.add_paragraph()
    p_auth.text = "Presented By: Prince Umrao  •  Raunak Kesharwani  •  Vibhu Kumar  •  Shubham Thakur"
    p_auth.font.size = Pt(13)
    p_auth.font.bold = True
    p_auth.font.color.rgb = CYAN_ACCENT
    p_auth.space_before = Pt(20)

    p_inst = tf_c.add_paragraph()
    p_inst.text = "B.Tech – Computer Science Engineering  |  Lamrin Tech Skills University"
    p_inst.font.size = Pt(12)
    p_inst.font.color.rgb = WHITE
    p_inst.space_before = Pt(4)

    output_path = "Retail_Pulse_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to: {os.path.abspath(output_path)}")

if __name__ == "__main__":
    create_presentation()
